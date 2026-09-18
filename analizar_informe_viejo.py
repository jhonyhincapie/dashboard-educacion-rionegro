# -*- coding: utf-8 -*-
"""Analiza estructura del informe viejo"""
from openpyxl import load_workbook

informe_viejo = r"C:\Users\jhincapie\Downloads\07-15 Seguimiento al presupuesto para claude.xlsx"

wb = load_workbook(informe_viejo)
ws = wb.active

print("ESTRUCTURA DEL INFORME VIEJO:")
print("=" * 100)
print(f"Hoja: {ws.title}, Filas: {ws.max_row}, Columnas: {ws.max_column}")
print("\n" + "=" * 100)
print("PRIMERAS 20 FILAS Y COLUMNAS A-H:")
print("=" * 100)

for r in range(1, min(21, ws.max_row + 1)):
    fila = []
    for c in range(1, min(9, ws.max_column + 1)):
        cell = ws.cell(r, c)
        val = str(cell.value or "")[:30].strip()
        fill = cell.fill
        fgColor = fill.fgColor if fill else None
        fila.append(f"{val:30s}")
    print(f"Fila {r:2d}: {'|'.join(fila)}")

# Analizar colores
print("\n\n" + "=" * 100)
print("CODIFICACIÓN DE COLORES:")
print("=" * 100)
colores = {}
for r in range(1, ws.max_row + 1):
    for c in range(1, min(9, ws.max_column + 1)):
        cell = ws.cell(r, c)
        if cell.fill and cell.fill.fgColor:
            color = cell.fill.fgColor.rgb if hasattr(cell.fill.fgColor, 'rgb') else str(cell.fill.fgColor.theme)
            val = str(cell.value or "")[:20]
            if color not in colores:
                colores[color] = {"ejemplo": val, "celdas": []}
            colores[color]["celdas"].append(f"{cell.coordinate}")

print(f"Total colores distintos encontrados: {len(colores)}")
for color, info in sorted(colores.items()):
    print(f"  Color {color}: '{info['ejemplo']}' ({len(info['celdas'])} celdas)")
    print(f"    Celdas: {', '.join(info['celdas'][:5])}")
