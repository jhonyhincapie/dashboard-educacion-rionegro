# -*- coding: utf-8 -*-
from openpyxl import load_workbook
import glob

# Buscar archivo CDP-CRP
archivos = glob.glob("entradas_contratos/Seguimiento CDP-CRP*.xlsx")
if not archivos:
    print("No se encontró el archivo CDP-CRP")
    exit(1)

file_path = archivos[0]
print(f"Leyendo: {file_path}\n")

wb = load_workbook(file_path, data_only=True)
ws = wb['2026']

sin_codigo = []
contrato = contratista = ''
for r in range(2, ws.max_row + 1):
    a = ws.cell(r, 1).value
    b = ws.cell(r, 2).value
    cod = ws.cell(r, 3).value
    vcdp = ws.cell(r, 5).value

    if a:
        contrato = str(a).strip()
    if b:
        contratista = str(b).strip()

    if cod in (None, '') and vcdp not in (None, '', 0):
        sin_codigo.append({
            'fila': r,
            'contrato': contrato,
            'contratista': contratista,
            'vcdp': str(vcdp).strip() if vcdp else '(sin)',
        })

print("FILAS SIN CODIGO PRESUPUESTAL (en hoja 2026):\n")
for i, info in enumerate(sin_codigo, 1):
    print(f"{i}. Fila {info['fila']}: Contrato {info['contrato']}")
    print(f"   Contratista: {info['contratista']}")
    print(f"   Valor_CDP: {info['vcdp']}")
    print()
