# -*- coding: utf-8 -*-
import os
import glob
from openpyxl import load_workbook

CDPCRP_DIR = "entradas_contratos"

archivos = glob.glob(os.path.join(CDPCRP_DIR, "Seguimiento CDP-CRP*.xlsx"))
print(f"Archivos encontrados: {archivos}")

if not archivos:
    print("No hay archivos")
    exit(1)

file_path = archivos[0]
print(f"Leyendo: {file_path}")

wb = load_workbook(file_path, data_only=True)
print(f"Hojas: {wb.sheetnames}")

ws = wb["2026"]
print(f"Max row: {ws.max_row}")

contratos = []
vigencias = set()
contrato = contratista = ""
sin_codigo_idx = 0

for r in range(2, ws.max_row + 1):
    a = ws.cell(r, 1).value
    b = ws.cell(r, 2).value
    cod = ws.cell(r, 3).value
    cdp = ws.cell(r, 5).value

    if a:
        contrato = str(a).strip()
    if b:
        contratista = str(b).strip()

    if not cdp or cdp in (None, "", 0):
        continue

    cons_ppt = ""
    incluir = False

    if cod and cod not in (None, ""):
        cod_str = str(cod).strip()
        if "1.19" in cod_str or "01.01.01" in cod_str:
            cons_ppt = cod_str
            incluir = True
    else:
        sin_codigo_idx += 1
        cons_ppt = "CDP_SIN_CODIGO_%d" % sin_codigo_idx
        incluir = True

    if not incluir or not cons_ppt:
        continue

    vigencias.add(cons_ppt)
    contratos.append({
        "nro_contrato": contrato,
        "contratista": contratista,
        "cons_ppt": cons_ppt,
    })

print(f"\nContratos encontrados: {len(contratos)}")
for c in contratos:
    print(f"  {c}")

print(f"\nVigencias a incluir: {vigencias}")
