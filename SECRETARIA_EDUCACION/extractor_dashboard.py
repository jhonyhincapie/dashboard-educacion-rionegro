#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Capa de DATOS del dashboard: lee la hoja SEGUIMIENTO del Excel y devuelve
la jerarquia Componente -> Subgrupo -> Rubro con valores exactos.
No calcula porcentajes ni brechas (eso es capa de calculo en el navegador)."""
import glob
import os
import re
import sys
from datetime import datetime

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
INFORMES = os.path.join(BASE, "informes")

# Mismo mapa componente -> color que usa extractor_educacion.py (MAPA_COLORES_DEFAULT)
COLOR_COMPONENTE = {
    "PRESTACION DEL SERVICIO": "FFFF00",
    "RECURSOS PROPIOS O LIBRE INVERSION": "CCFFFF",
    "RENDIMIENTOS O RECURSOS DEL BALANCE DE PRESTACION DEL SERVICIO": "FFFFCC",
    "SGP - PGN - PAE": "00CCFF",
    "CALIDAD MATRICULA": "99CC00",
    "PRIMERA INFANCIA": "FFCC99",
    "CALIDAD GRATUIDAD": "FFCC00",
    "INFRAESTRUCTURA": "FF99CC",
    "VIGENCIAS ANTERIORES": "C0C0C0",
}
# Columnas de SEGUIMIENTO (fila 1 = encabezados)
COL = dict(comp=1, nat=2, codigo=3, cons=4, rubro=5, fondo=6, fuente=7,
           definitivo=8, reservado=9, comprometido=10, obligado=11, pagado=12,
           contratos=18, responsable=19, subgrupo=21)
ETAPAS = ("definitivo", "reservado", "comprometido", "obligado", "pagado")
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def excel_mas_reciente():
    """Elige el informe con corte (dd_mm_yyyy en el nombre) mas reciente."""
    def corte(p):
        m = re.search(r"(\d{2})_(\d{2})_(\d{4})", os.path.basename(p))
        return (m.group(3), m.group(2), m.group(1)) if m else ("", "", "")
    archivos = [p for p in glob.glob(os.path.join(INFORMES, "Seguimiento Presupuestal*.xlsx"))
                if not os.path.basename(p).startswith("~$")]
    return max(archivos, key=corte) if archivos else None


def _num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def extraer(excel_path=None):
    excel_path = excel_path or excel_mas_reciente()
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except PermissionError:  # archivo abierto en Excel: leer una copia temporal
        import shutil, tempfile
        tmp = os.path.join(tempfile.gettempdir(), "_seg_tmp.xlsx")
        shutil.copyfile(excel_path, tmp)
        wb = openpyxl.load_workbook(tmp, data_only=True)
    ws = wb["SEGUIMIENTO"]

    comps, orden = {}, []
    total_excel = None
    subtotales_excel = {}
    for r in range(2, ws.max_row + 1):
        nombre = ws.cell(r, COL["comp"]).value
        if not nombre:
            continue
        nombre = str(nombre).strip()
        vals = {e: _num(ws.cell(r, COL[e]).value) for e in ETAPAS}
        if nombre.upper() == "TOTAL GENERAL":
            if ws.cell(r, COL["definitivo"]).value is not None:
                total_excel = vals
            continue
        if nombre.upper().startswith("SUBTOTAL"):
            if ws.cell(r, COL["definitivo"]).value is not None:
                subtotales_excel[nombre[len("SUBTOTAL"):].strip()] = vals
            continue
        sub = (ws.cell(r, COL["subgrupo"]).value or nombre)
        sub = str(sub).strip()
        if nombre not in comps:
            comps[nombre] = {"nombre": nombre,
                             "color": COLOR_COMPONENTE.get(nombre, "FFFFFF"),
                             "subgrupos": {}}
            orden.append(nombre)
        rubro = {"cons_ppt": ws.cell(r, COL["cons"]).value,
                 "codigo": ws.cell(r, COL["codigo"]).value,
                 "nombre": ws.cell(r, COL["rubro"]).value,
                 "fondo": ws.cell(r, COL["fondo"]).value,
                 "fuente": ws.cell(r, COL["fuente"]).value,
                 "responsable": ws.cell(r, COL["responsable"]).value,
                 **vals}
        comps[nombre]["subgrupos"].setdefault(sub, []).append(rubro)

    componentes = []
    for n in orden:
        c = comps[n]
        c["subgrupos"] = [{"nombre": s, "rubros": rs} for s, rs in c["subgrupos"].items()]
        componentes.append(c)

    # Validacion: suma de rubros vs TOTAL GENERAL y SUBTOTALES del Excel
    suma = {e: 0.0 for e in ETAPAS}
    diffs = []
    for c in componentes:
        sc = {e: 0.0 for e in ETAPAS}
        for s in c["subgrupos"]:
            for rb in s["rubros"]:
                for e in ETAPAS:
                    sc[e] += rb[e]
        for e in ETAPAS:
            suma[e] += sc[e]
        ref = subtotales_excel.get(c["nombre"])
        if ref:
            for e in ETAPAS:
                if abs(sc[e] - ref[e]) > 1:
                    diffs.append({"dato": f"{c['nombre']} / {e}", "fuente": "SUBTOTAL en SEGUIMIENTO",
                                  "calculado": sc[e], "excel": ref[e], "diferencia": sc[e] - ref[e]})
    if total_excel:
        for e in ETAPAS:
            if abs(suma[e] - total_excel[e]) > 1:
                diffs.append({"dato": f"TOTAL GENERAL / {e}", "fuente": "TOTAL GENERAL en SEGUIMIENTO",
                              "calculado": suma[e], "excel": total_excel[e], "diferencia": suma[e] - total_excel[e]})

    m = re.search(r"(\d{2})_(\d{2})_(\d{4})", os.path.basename(excel_path))
    corte = (f"{int(m.group(1))} de {MESES[int(m.group(2)) - 1]} de {m.group(3)}" if m else "")
    return {
        "corte": corte,
        "archivo": os.path.basename(excel_path),
        "generado": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        "total_excel": total_excel,
        "validacion": {"ok": (not diffs) if (total_excel or subtotales_excel) else None,
                       "diferencias": diffs, "rubros": sum(
            len(s["rubros"]) for c in componentes for s in c["subgrupos"])},
        "componentes": componentes,
    }


if __name__ == "__main__":
    import json
    d = extraer(sys.argv[1] if len(sys.argv) > 1 else None)
    print(json.dumps({k: d[k] for k in ("corte", "archivo", "total_excel", "validacion")},
                     ensure_ascii=False, indent=2))
